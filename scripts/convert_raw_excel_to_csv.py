"""Convert raw Excel workbooks with .csv suffixes into UTF-8 CSV files.

The files under data/raw are treated as immutable source data. Converted files
are written under data/staging with lowercase, hyphenated filenames.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import tempfile
from collections.abc import Iterable
from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_excel_workbook(path: Path) -> bool:
    try:
        with ZipFile(path) as archive:
            names = set(archive.namelist())
    except BadZipFile:
        return False
    return "xl/workbook.xml" in names and "[Content_Types].xml" in names


def csv_value(value: Any) -> Any:
    if isinstance(value, datetime):
        if value.time() == time.min:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if value is None:
        return ""
    return value


def trimmed_rows(rows: Iterable[tuple[Any, ...]], column_count: int) -> Iterable[list[Any]]:
    for row in rows:
        yield [csv_value(value) for value in row[:column_count]]


def normalized_relative_path(path: Path) -> Path:
    stem = re.sub(r"[^a-z0-9]+", "-", path.stem.lower()).strip("-")
    return path.with_name(stem + path.suffix.lower())


def convert(source: Path, destination: Path) -> dict[str, Any]:
    destination.parent.mkdir(parents=True, exist_ok=True)

    with source.open("rb") as source_stream:
        workbook = load_workbook(source_stream, read_only=True, data_only=True)
        if len(workbook.sheetnames) != 1:
            raise ValueError(
                f"Expected one worksheet in {source}, found {len(workbook.sheetnames)}"
            )
        worksheet = workbook.active
        assert worksheet is not None
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError(f"Workbook is empty: {source}") from exc

        nonempty_header_indexes = [
            index for index, value in enumerate(header) if value is not None
        ]
        if not nonempty_header_indexes:
            raise ValueError(f"Workbook has no header: {source}")
        column_count = max(nonempty_header_indexes) + 1
        header_values = ["source_row"] + [
            csv_value(value) for value in header[:column_count]
        ]

        file_descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent
        )
        row_count = 1
        try:
            with os.fdopen(file_descriptor, "w", encoding="utf-8", newline="") as output:
                writer = csv.writer(output, lineterminator="\n")
                writer.writerow(header_values)
                for source_row, row in enumerate(
                    trimmed_rows(rows, column_count), start=2
                ):
                    writer.writerow([source_row, *row])
                    row_count += 1
            os.replace(temporary_name, destination)
        except BaseException:
            Path(temporary_name).unlink(missing_ok=True)
            raise
        finally:
            workbook.close()

    return {
        "source": str(source),
        "destination": str(destination),
        "worksheet": worksheet.title,
        "rows_including_header": row_count,
        "data_rows": row_count - 1,
        "columns": len(header_values),
        "header": header_values,
        "source_bytes": source.stat().st_size,
        "source_sha256": sha256(source),
        "destination_bytes": destination.stat().st_size,
        "destination_sha256": sha256(destination),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=Path("data/raw"))
    parser.add_argument("--staging-dir", type=Path, default=Path("data/staging"))
    parser.add_argument(
        "--manifest", type=Path, default=Path("data/staging/conversion-manifest.json")
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    sources = sorted(
        path for path in args.raw_dir.rglob("*.csv") if is_excel_workbook(path)
    )
    if not sources:
        raise SystemExit(f"No Excel workbooks with .csv suffixes found under {args.raw_dir}")

    records = []
    for source in sources:
        relative_path = normalized_relative_path(source.relative_to(args.raw_dir))
        destination = args.staging_dir / relative_path
        record = convert(source, destination)
        records.append(record)
        print(
            f"converted {source} -> {destination} "
            f"({record['data_rows']} data rows, {record['columns']} columns)",
            flush=True,
        )

    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    manifest = {
        "conversion": "Excel Open XML workbook with misleading .csv suffix to UTF-8 CSV",
        "raw_data_modified": False,
        "file_count": len(records),
        "files": records,
    }
    args.manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"wrote manifest {args.manifest}")


if __name__ == "__main__":
    main()
